#!/usr/bin/env python3
"""
Unit tests for Step 2: Dynamic Excel Workbook & Requirement Analyzer.
Tests PRD Sections 16, 17, 18, 19 adherence:
- Dynamic header detection
- Multi-level composite sub-headers
- Merged cell resolution
- Section hierarchy & requirement extraction
- Target fillable slots mapping
- Zero Any typing compliance
"""

from pathlib import Path

import openpyxl
import pytest
from step2_parse_excel import (
    classify_column_header,
    detect_headers,
    export_questions_list,
    get_merged_cell_value,
    parse_excel_workbook,
)


@pytest.fixture
def sample_rfp_excel(tmp_path: Path) -> Path:
    """Generates a realistic RFP Excel workbook with sections, merged cells, and composite subheaders."""
    excel_path = tmp_path / "test_rfp_spec.xlsx"
    wb = openpyxl.Workbook()

    # Sheet 1: Server Requirements (Multi-level headers)
    ws1 = wb.active
    ws1.title = "Server Specs"

    # Row 1: Document title banner
    ws1.merge_cells("A1:E1")
    ws1["A1"] = "SECTION 3: TECHNICAL SERVER SPECIFICATIONS"

    # Row 2: Header Level 1
    ws1["A2"] = "S.No"
    ws1["B2"] = "Technical Requirement"
    ws1.merge_cells("C2:E2")
    ws1["C2"] = "Bidder Evaluation (Techaccess)"

    # Row 3: Subheader Level 2
    ws1["A3"] = "#"
    ws1["B3"] = "Specification Details"
    ws1["C3"] = "Compliance (Yes/No)"
    ws1["D3"] = "Marks Obtained"
    ws1["E3"] = "Remarks & Evidence"

    # Row 4: Section Banner
    ws1.merge_cells("A4:E4")
    ws1["A4"] = "3.1 Compute & Processor Specifications"

    # Row 5: Requirement 1
    ws1["A5"] = 1
    ws1["B5"] = "Dual Intel Xeon 6th Gen Scalable Processors (48-Core min)"
    ws1["C5"] = None  # Empty slot to be filled
    ws1["D5"] = None  # Empty slot to be filled
    ws1["E5"] = None  # Empty slot to be filled

    # Row 6: Requirement 2
    ws1["A6"] = 2
    ws1["B6"] = "Minimum 512GB DDR5 Registered ECC Memory (16x32GB)"
    ws1["C6"] = None
    ws1["D6"] = None
    ws1["E6"] = None

    # Row 7: Section Banner
    ws1.merge_cells("A7:E7")
    ws1["A7"] = "3.2 Storage & Controller"

    # Row 8: Requirement 3
    ws1["A8"] = 3
    ws1["B8"] = "All-Flash Storage delivering 350 TB RAW Capacity with NVMe SSDs"
    ws1["C8"] = None
    ws1["D8"] = None
    ws1["E8"] = None

    # Sheet 2: Simple Flat Requirements
    ws2 = wb.create_sheet("Network Specs")
    ws2["A1"] = "Item No"
    ws2["B1"] = "Requirement Description"
    ws2["C1"] = "Complied / Not Complied"
    ws2["D1"] = "Bidder Notes"

    ws2["A2"] = 1
    ws2["B2"] = "8x 32G Fibre Channel Host Ports"
    ws2["C2"] = None
    ws2["D2"] = None

    wb.save(excel_path)
    wb.close()
    return excel_path


def test_classify_column_header() -> None:
    """Validates classification of various industry column headers."""
    assert classify_column_header("S.No") == "index"
    assert classify_column_header("Sr No") == "index"
    assert classify_column_header("Technical Requirement") == "requirement"
    assert classify_column_header("Specification Details") == "requirement"
    assert classify_column_header("Compliance (Yes/No)") == "compliance"
    assert classify_column_header("Complied / Not Complied") == "compliance"
    assert classify_column_header("Total Marks") == "total_marks"
    assert classify_column_header("Marks Obtained") == "marks"
    assert classify_column_header("Remarks & Evidence") == "remarks"
    assert classify_column_header("Offered Specification") == "answer"


def test_merged_cell_resolution(sample_rfp_excel: Path) -> None:
    """Ensures get_merged_cell_value retrieves master cell value across merged spans."""
    wb = openpyxl.load_workbook(sample_rfp_excel, data_only=False)
    ws = wb["Server Specs"]

    # In A1:E1, cells B1..E1 should resolve to A1's value
    assert get_merged_cell_value(ws, 1, 1) == "SECTION 3: TECHNICAL SERVER SPECIFICATIONS"
    assert get_merged_cell_value(ws, 1, 3) == "SECTION 3: TECHNICAL SERVER SPECIFICATIONS"
    assert get_merged_cell_value(ws, 1, 5) == "SECTION 3: TECHNICAL SERVER SPECIFICATIONS"

    wb.close()


def test_detect_multi_level_headers(sample_rfp_excel: Path) -> None:
    """Validates composite header resolution for multi-level tables."""
    wb = openpyxl.load_workbook(sample_rfp_excel, data_only=False)
    ws = wb["Server Specs"]

    header_row, data_start, headers = detect_headers(ws)
    assert header_row == 2
    assert data_start == 4  # Row 2 (main) + Row 3 (sub) -> Data starts at row 4
    assert "S.No" in headers[1]
    assert "Requirement" in headers[2]
    assert "Compliance" in headers[3]
    assert "Marks" in headers[4]
    assert "Remarks" in headers[5]

    wb.close()


def test_parse_excel_workbook(sample_rfp_excel: Path) -> None:
    """Validates complete WorkbookAnalysis extraction and slot mapping."""
    analysis = parse_excel_workbook(sample_rfp_excel)

    assert analysis["workbook_id"] == "test_rfp_spec"
    assert len(analysis["sheets"]) == 2
    assert analysis["total_requirements"] == 4  # 3 in Sheet 1, 1 in Sheet 2

    # Verify Sheet 1
    s1 = analysis["sheets"][0]
    assert s1["sheet_name"] == "Server Specs"
    assert len(s1["requirements"]) == 3

    # Check Requirement 1
    req1 = s1["requirements"][0]
    assert req1["requirement_id"] == "REQ-S01-R001"
    assert "Intel Xeon" in req1["requirement_text"]
    assert req1["row_number"] == 5
    assert req1["source_cell"] == "B5"
    assert req1["section"] == "3.1 Compute & Processor Specifications"

    # Verify Target Slots
    slots = req1["target_slots"]
    assert "compliance" in slots
    assert slots["compliance"]["cell_coordinate"] == "C5"
    assert "marks" in slots
    assert slots["marks"]["cell_coordinate"] == "D5"
    assert "remarks" in slots
    assert slots["remarks"]["cell_coordinate"] == "E5"

    # Check Requirement 3 (Section change)
    req3 = s1["requirements"][2]
    assert req3["section"] == "3.2 Storage & Controller"
    assert "350 TB" in req3["requirement_text"]

    # Verify Sheet 2
    s2 = analysis["sheets"][1]
    assert s2["sheet_name"] == "Network Specs"
    assert len(s2["requirements"]) == 1
    assert "Fibre Channel" in s2["requirements"][0]["requirement_text"]


def test_export_questions_list(sample_rfp_excel: Path) -> None:
    """Validates export of flattened question/requirement list for reverse probing."""
    analysis = parse_excel_workbook(sample_rfp_excel)
    questions = export_questions_list(analysis)

    assert len(questions) == 4
    texts = [q["requirement_text"] for q in questions]
    assert any("Xeon" in t for t in texts)
    assert any("512GB" in t for t in texts)
    assert any("350 TB" in t for t in texts)
    assert any("Fibre Channel" in t for t in texts)
