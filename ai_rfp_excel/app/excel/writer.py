from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ai_rfp_excel.app.config import settings
from ai_rfp_excel.app.excel.models import (
    PopulationResult,
    WorkbookAnalysis,
)
from ai_rfp_excel.app.excel.validator import ExcelValidator
from ai_rfp_excel.app.logging import get_logger
from ai_rfp_excel.app.matching.models import ComplianceDecision, ComplianceState

logger = get_logger("excel.writer")

# Soft highlight colors for compliance states (hex fill colors)
STATUS_FILLS = {
    ComplianceState.COMPLIANT: PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid"),
    ComplianceState.NON_COMPLIANT: PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    ComplianceState.PARTIALLY_COMPLIANT: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    ComplianceState.AMBIGUOUS: PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    ComplianceState.NOT_FOUND: PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}

STATUS_LABELS = {
    ComplianceState.COMPLIANT: "Compliant",
    ComplianceState.NON_COMPLIANT: "Non-Compliant",
    ComplianceState.PARTIALLY_COMPLIANT: "Partially Compliant",
    ComplianceState.AMBIGUOUS: "Ambiguous",
    ComplianceState.NOT_FOUND: "Not Found",
}


class ExcelWriter:
    """Populates RFP response workbooks with compliance decisions and evidence."""

    def __init__(self, output_dir: str | Path | None = None) -> None:
        self.output_dir = Path(output_dir) if output_dir else Path(settings.OUTPUT_DIR)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.validator = ExcelValidator()

    def _safe_set_cell_value(
        self,
        ws: Worksheet,
        cell_coord: str,
        value: Any,
        fill: PatternFill | None = None,
        only_if_empty: bool = False,
    ) -> bool:
        """Safely write to normal cells and merged cells with auto text-wrapping and clean bounds."""
        if isinstance(value, str):
            # Sanitize illegal xml chars
            clean_str = "".join(
                c
                for c in value
                if (
                    ord(c) in (0x9, 0xA, 0xD)
                    or (0x20 <= ord(c) <= 0xD7FF)
                    or (0xE000 <= ord(c) <= 0xFFFD)
                    or (0x10000 <= ord(c) <= 0x10FFFF)
                )
            ).strip()
            # Truncate oversized raw text dumps to prevent Excel distortion
            if len(clean_str) > 250:
                clean_str = clean_str[:247] + "..."
            value = clean_str

        cell = ws[cell_coord]
        target_cell = cell

        if isinstance(cell, MergedCell):
            for rng in ws.merged_cells.ranges:
                if cell.coordinate in rng:
                    target_cell = ws.cell(row=rng.min_row, column=rng.min_col)
                    break

        if only_if_empty and target_cell.value is not None and str(target_cell.value).strip():
            return False

        target_cell.value = value
        target_cell.alignment = Alignment(wrap_text=True, vertical="top")
        if fill:
            target_cell.fill = fill

        return True

    def populate_workbook(
        self,
        template_path: str | Path,
        analysis: WorkbookAnalysis,
        decisions: list[ComplianceDecision],
        create_summary: bool = True,
    ) -> PopulationResult:
        """Populate template workbook with compliance results, preserving formatting and formulas."""
        template_file = Path(template_path)
        if not template_file.exists():
            raise FileNotFoundError(f"Template workbook not found: {template_path}")

        # Load original workbook preserving formulas (data_only=False)
        wb = openpyxl.load_workbook(template_file, data_only=False)

        # Index decisions for lookup
        decision_by_id = {d.requirement_id: d for d in decisions if d.requirement_id}
        decision_by_text = {d.requirement_text.strip().lower(): d for d in decisions}

        total_populated = 0
        all_evaluated_decisions: list[ComplianceDecision] = []

        # Iterate through sheets and requirements
        for sheet_analysis in analysis.sheets:
            if sheet_analysis.sheet_name not in wb.sheetnames:
                continue

            ws = wb[sheet_analysis.sheet_name]

            for req in sheet_analysis.requirements:
                # Find matching decision
                dec = decision_by_id.get(req.requirement_id)
                if not dec:
                    dec = decision_by_text.get(req.requirement_text.strip().lower())

                if not dec:
                    continue

                all_evaluated_decisions.append(dec)
                status_str = STATUS_LABELS.get(dec.state, dec.state.value)
                fill = STATUS_FILLS.get(dec.state)

                # Format evidence text for remarks (concise citation only)
                remarks_parts: list[str] = []
                if dec.evidence:
                    for ev in dec.evidence:
                        if ev.citation:
                            remarks_parts.append(ev.citation)
                        elif ev.reasoning:
                            remarks_parts.append(ev.reasoning[:120])
                elif dec.reasoning:
                    remarks_parts.append(dec.reasoning[:120])

                remarks_text = " | ".join(dict.fromkeys(remarks_parts)) if remarks_parts else "Specification verified."

                # 1. Populate compliance cells
                for _, cell_coord in req.compliance_cells.items():
                    if self._safe_set_cell_value(ws, cell_coord, status_str, fill):
                        total_populated += 1

                # 2. Populate answer / offered specification / proposed value cells
                matched_val = dec.matched_value or (dec.evidence[0].value if dec.evidence else "")
                if matched_val:
                    for _, cell_coord in req.answer_cells.items():
                        if self._safe_set_cell_value(ws, cell_coord, matched_val):
                            total_populated += 1

                    for _, cell_coord in req.proposed_cells.items():
                        if self._safe_set_cell_value(ws, cell_coord, matched_val):
                            total_populated += 1

                    for _, cell_coord in req.offered_spec_cells.items():
                        if self._safe_set_cell_value(ws, cell_coord, matched_val):
                            total_populated += 1

                    for _, cell_coord in req.vendor_cells.items():
                        # Only fill vendor cell if not already populated
                        if (
                            cell_coord not in req.compliance_cells.values()
                            and cell_coord not in req.offered_spec_cells.values()
                            and cell_coord not in req.answer_cells.values()
                            and cell_coord not in req.proposed_cells.values()
                        ):
                            if self._safe_set_cell_value(ws, cell_coord, matched_val):
                                total_populated += 1

                # 3. Populate remarks cells
                for _, cell_coord in req.remarks_cells.items():
                    if self._safe_set_cell_value(ws, cell_coord, remarks_text):
                        total_populated += 1

                # 4. Populate total marks cells ONLY IF EMPTY
                is_preference = "preference" in req.requirement_text.lower() or "preferred" in req.requirement_text.lower()
                total_marks_val = "10" if is_preference else "Mandatory"
                for _, cell_coord in req.total_marks_cells.items():
                    if self._safe_set_cell_value(ws, cell_coord, total_marks_val, only_if_empty=True):
                        total_populated += 1

                # 5. Populate marks obtained cells
                if dec.state == ComplianceState.COMPLIANT:
                    marks_val = "10"
                elif dec.state == ComplianceState.PARTIALLY_COMPLIANT:
                    marks_val = "5"
                else:
                    marks_val = "0"

                for _, cell_coord in req.marks_cells.items():
                    if self._safe_set_cell_value(ws, cell_coord, marks_val):
                        total_populated += 1

            # Auto-fit column widths for readable, unclipped layout
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    if cell.value:
                        lines = str(cell.value).split("\n")
                        max_len = max(max_len, max(len(line_str) for line_str in lines))
                if max_len > 0:
                    col_letter = get_column_letter(col[0].column)
                    ws.column_dimensions[col_letter].width = max(14, min(max_len + 3, 50))

        # Add Compliance Summary Sheet if requested
        if create_summary:
            self._create_summary_sheet(wb, analysis, all_evaluated_decisions or decisions)

        # File naming: {original_name}_populated_{timestamp}.xlsx
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        stem = template_file.stem
        populated_filename = f"{stem}_populated_{timestamp}.xlsx"
        output_file_path = self.output_dir / populated_filename

        wb.save(output_file_path)
        logger.info("Populated workbook saved", path=str(output_file_path), populated_cells=total_populated)

        # Run automated validation
        validation_report = self.validator.validate(
            populated_path=output_file_path,
            original_analysis=analysis,
            decisions=all_evaluated_decisions or decisions,
        )

        return PopulationResult(
            output_file_path=str(output_file_path),
            filename=populated_filename,
            total_populated_cells=total_populated,
            summary_sheet_created=True,
            validation_report=validation_report,
        )

    def _create_summary_sheet(
        self,
        wb: openpyxl.Workbook,
        analysis: WorkbookAnalysis,
        decisions: list[ComplianceDecision],
    ) -> None:
        """Create Executive Compliance Dashboard tab at the beginning of the workbook."""
        summary_title = "Compliance Summary"
        if summary_title in wb.sheetnames:
            del wb[summary_title]

        ws = wb.create_sheet(title=summary_title, index=0)
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws.merge_cells("A1:G2")
        title_cell = ws["A1"]
        title_cell.value = "EXECUTIVE RFP COMPLIANCE DASHBOARD"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        title_cell.fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")

        # Aggregate Statistics
        total_reqs = len(decisions)
        compliant_count = sum(1 for d in decisions if d.state == ComplianceState.COMPLIANT)
        non_compliant_count = sum(1 for d in decisions if d.state == ComplianceState.NON_COMPLIANT)
        partial_count = sum(1 for d in decisions if d.state == ComplianceState.PARTIALLY_COMPLIANT)
        ambiguous_count = sum(1 for d in decisions if d.state == ComplianceState.AMBIGUOUS)
        not_found_count = sum(1 for d in decisions if d.state == ComplianceState.NOT_FOUND)

        comp_pct = (compliant_count / total_reqs * 100) if total_reqs > 0 else 0.0

        # KPI Summary Table
        headers = ["Compliance Category", "Count", "Percentage"]
        for col_num, h_text in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_num)
            cell.value = h_text
            cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        stats_data = [
            ("Fully Compliant", compliant_count, compliant_count / total_reqs if total_reqs > 0 else 0, "E2F0D9"),
            ("Partially Compliant", partial_count, partial_count / total_reqs if total_reqs > 0 else 0, "FFF2CC"),
            ("Non-Compliant", non_compliant_count, non_compliant_count / total_reqs if total_reqs > 0 else 0, "FCE4D6"),
            ("Ambiguous / Under Review", ambiguous_count, ambiguous_count / total_reqs if total_reqs > 0 else 0, "FFF2CC"),
            ("Not Found in Reference", not_found_count, not_found_count / total_reqs if total_reqs > 0 else 0, "F2F2F2"),
        ]

        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )

        for i, (label, count, pct, color_hex) in enumerate(stats_data, 5):
            c1 = ws.cell(row=i, column=1, value=label)
            c2 = ws.cell(row=i, column=2, value=count)
            c3 = ws.cell(row=i, column=3, value=pct)

            c1.fill = PatternFill(start_color=color_hex, end_color=color_hex, fill_type="solid")
            c2.alignment = Alignment(horizontal="center")
            c3.alignment = Alignment(horizontal="center")
            c3.number_format = "0.0%"

            for c in (c1, c2, c3):
                c.border = thin_border
                c.font = Font(name="Calibri", size=11)

        # Total Row
        ws.cell(row=10, column=1, value="Total Evaluated").font = Font(bold=True)
        ws.cell(row=10, column=2, value=total_reqs).font = Font(bold=True)
        ws.cell(row=10, column=3, value=1.0).font = Font(bold=True)
        ws.cell(row=10, column=3).number_format = "0.0%"

        # KPI Score Card
        ws.merge_cells("E4:G6")
        kpi_cell = ws["E4"]
        kpi_cell.value = f"COMPLIANCE SCORE\n{comp_pct:.1f}%"
        kpi_cell.font = Font(name="Calibri", size=16, bold=True, color="1F4E79")
        kpi_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        kpi_cell.fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        kpi_cell.border = thin_border

        # Embed Native Excel Charts
        pie = PieChart()
        pie.title = "Compliance Distribution"
        labels = Reference(ws, min_col=1, min_row=5, max_row=9)
        data = Reference(ws, min_col=2, min_row=4, max_row=9)
        pie.add_data(data, titles_from_data=True)
        pie.set_categories(labels)
        pie.width = 14
        pie.height = 7
        ws.add_chart(pie, "A12")

        bar = BarChart()
        bar.title = "Evaluation Breakdown"
        bar.style = 10
        bar.y_axis.title = "Count"
        bar.x_axis.title = "Status"
        bar.add_data(data, titles_from_data=True)
        bar.set_categories(labels)
        bar.legend = None
        bar.width = 14
        bar.height = 7
        ws.add_chart(bar, "E12")

        # Column widths for Summary tab
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 14
        ws.column_dimensions["D"].width = 5
        ws.column_dimensions["E"].width = 16
        ws.column_dimensions["F"].width = 16
        ws.column_dimensions["G"].width = 16
