"""
Step 5: "The In-Place Excel Populator & Styler" — Final Output Generation.
Consumes Step 4 (compliance_decisions.json) and updates the original template workbook:
1. In-place workbook modification using openpyxl (Rule 9).
2. Formula Guard: strictly protects existing Excel formulas.
3. Rule 7 Visual Auditing: soft yellow fill (FFF2CC) on needs_review=True cells.
4. Generates an inspectable JSON artifact: population_manifest.json (Rule 8).

Strict typing only — zero Any (Rule 4). 100% offline (Rule 10).
"""

from __future__ import annotations

import copy
import datetime
import json
import logging
from pathlib import Path
from typing import Final, TypedDict, cast

import openpyxl
from openpyxl.cell.cell import Cell
from openpyxl.styles import PatternFill
from openpyxl.utils import coordinate_to_tuple
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")
logger = logging.getLogger("Step5PopulateExcel")

# Styling Tokens per PRD Rule 7 & Rule 9
YELLOW_HIGHLIGHT_HEX: Final[str] = "FFF2CC"
YELLOW_FILL: Final[PatternFill] = PatternFill(
    start_color=YELLOW_HIGHLIGHT_HEX,
    end_color=YELLOW_HIGHLIGHT_HEX,
    fill_type="solid",
)
EXCEL_MAX_ROWS: Final[int] = 1_048_576
EXCEL_MAX_COLS: Final[int] = 16_384


class UpdatedCellRecord(TypedDict):
    sheet_name: str
    cell_coordinate: str
    slot_type: str
    value: str
    needs_review: bool
    highlighted: bool
    is_merged: bool


class SheetSummary(TypedDict):
    sheet_name: str
    cells_updated: int
    cells_highlighted: int
    formulas_preserved: int
    collisions_detected: int


class PopulationManifest(TypedDict):
    document_id: str
    workbook_id: str
    source_file: str
    output_excel_file: str
    generated_at: str
    total_requirements: int
    total_cells_updated: int
    unique_cells_updated: int
    total_cells_highlighted: int
    total_formulas_preserved: int
    total_collisions_detected: int
    sheet_summaries: list[SheetSummary]
    updated_cells: list[UpdatedCellRecord]


def normalize_and_validate_coordinate(coord: str) -> str | None:
    """Clean and validate coordinate, stripping $ and checking Excel grid limits."""
    clean = coord.replace("$", "").strip().upper()
    if not clean:
        return None
    try:
        row, col = coordinate_to_tuple(clean)
        if 1 <= row <= EXCEL_MAX_ROWS and 1 <= col <= EXCEL_MAX_COLS:
            return clean
        return None
    except ValueError:
        return None


def find_sheet_case_insensitive(wb: Workbook, requested_name: str) -> tuple[Worksheet, str]:
    """Locate worksheet case-insensitively with whitespace tolerance."""
    req_clean = requested_name.strip().lower()
    for name in wb.sheetnames:
        if name.strip().lower() == req_clean:
            return cast(Worksheet, wb[name]), name

    # Fallback to active sheet or first available sheet
    active_ws = wb.active
    ws: Worksheet = cast(Worksheet, active_ws if active_ws is not None else wb[wb.sheetnames[0]])
    logger.warning("Sheet '%s' not found in workbook; falling back to '%s'.", requested_name, ws.title)
    return ws, ws.title


def resolve_target_cell(ws: Worksheet, coord: str) -> tuple[Cell, list[Cell]]:
    """
    Resolve cell coordinate to a writable anchor Cell.
    If coordinate is part of a MergedCellRange, returns top-left anchor
    and all cells in the range for uniform styling.
    """
    for merged_range in ws.merged_cells.ranges:
        if coord in merged_range:
            anchor = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
            all_cells = [
                ws.cell(row=r, column=c)
                for r in range(merged_range.min_row, merged_range.max_row + 1)
                for c in range(merged_range.min_col, merged_range.max_col + 1)
            ]
            return anchor, all_cells

    cell = ws[coord]
    return cell, [cell]


def is_formula_cell(cell: Cell) -> bool:
    """Check if an openpyxl cell contains an Excel formula."""
    if getattr(cell, "data_type", None) == "f":
        return True
    val = getattr(cell, "value", None)
    return isinstance(val, str) and val.strip().startswith("=")


def apply_cell_update(
    anchor_cell: Cell,
    value: str,
    needs_review: bool,
    all_cells: list[Cell] | None = None,
) -> tuple[bool, bool]:
    """
    Safely write value to cell and apply yellow highlight if needs_review is True.
    Preserves formulas (Rule 9) and prevents formula injection / #NAME? errors.
    Returns (updated, highlighted).
    """
    if is_formula_cell(anchor_cell):
        logger.warning("Protected formula cell %s detected; skipping write.", anchor_cell.coordinate)
        return False, False

    # Rule 7: Zero hallucination for missing requirements
    if value == "NOT_SPECIFIED" or not value.strip():
        value = "NOT_SPECIFIED"
        needs_review = True

    anchor_cell.value = value

    # Prevent accidental formula evaluation for spec text starting with =, +, -, @
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        anchor_cell.data_type = "s"

    # Multiline text wrap
    if "\n" in value:
        alignment_copy = copy.copy(anchor_cell.alignment)
        alignment_copy.wrap_text = True
        anchor_cell.alignment = alignment_copy

    highlighted = False
    if needs_review:
        target_cells = all_cells if all_cells is not None else [anchor_cell]
        for c in target_cells:
            c.fill = YELLOW_FILL
        highlighted = True

    return True, highlighted


def populate_workbook_in_place(
    decisions_path: Path,
    original_excel_path: Path,
    output_excel_path: Path,
    manifest_output_path: Path | None = None,
) -> PopulationManifest:
    """Load template workbook, apply Step 4 slot assignments, and save populated output."""
    if not decisions_path.exists():
        raise FileNotFoundError(f"Decisions artifact not found: {decisions_path}")
    if not original_excel_path.exists():
        raise FileNotFoundError(f"Original Excel template not found: {original_excel_path}")

    with open(decisions_path, "r", encoding="utf-8") as f:
        loaded_json = json.load(f)

    decisions_doc: dict[str, object] = loaded_json if isinstance(loaded_json, dict) else {}

    doc_id = str(decisions_doc.get("document_id", "DOC-UNKNOWN"))
    wb_id = str(decisions_doc.get("workbook_id", "WB-UNKNOWN"))
    items_raw = decisions_doc.get("items", [])
    items = [it for it in items_raw if isinstance(it, dict)] if isinstance(items_raw, list) else []

    wb: Workbook = openpyxl.load_workbook(original_excel_path, data_only=False)
    default_sheet_title = wb.active.title if wb.active is not None else wb.sheetnames[0]

    updated_records: list[UpdatedCellRecord] = []
    seen_coords: dict[str, set[str]] = {}
    total_formulas_preserved = 0
    total_collisions = 0
    sheet_stats: dict[str, dict[str, int]] = {}

    for item in items:
        raw_sheet = str(item.get("sheet_name", default_sheet_title))
        ws, sheet_name = find_sheet_case_insensitive(wb, raw_sheet)

        if sheet_name not in sheet_stats:
            sheet_stats[sheet_name] = {"updated": 0, "highlighted": 0, "formulas": 0, "collisions": 0}
            seen_coords[sheet_name] = set()

        slots_raw = item.get("slot_assignments", {})
        if not isinstance(slots_raw, dict):
            continue

        for slot_key, slot_obj in slots_raw.items():
            if not isinstance(slot_obj, dict):
                continue

            raw_coord = str(slot_obj.get("cell_coordinate", "")).strip()
            cell_coord = normalize_and_validate_coordinate(raw_coord)
            if not cell_coord:
                logger.warning("Invalid cell coordinate '%s'; skipping slot.", raw_coord)
                continue

            raw_val = slot_obj.get("value")
            if raw_val is None or str(raw_val).strip() in ("", "None"):
                val = "NOT_SPECIFIED"
                needs_review = True
            else:
                val = str(raw_val).strip()
                needs_review = bool(slot_obj.get("needs_review", False))

            if val == "NOT_SPECIFIED":
                needs_review = True

            slot_type = str(slot_obj.get("slot_type", slot_key))

            # Collision detection
            if cell_coord in seen_coords[sheet_name]:
                logger.warning("Coordinate collision: cell '%s' in sheet '%s' targeted multiple times.", cell_coord, sheet_name)
                total_collisions += 1
                sheet_stats[sheet_name]["collisions"] += 1
            else:
                seen_coords[sheet_name].add(cell_coord)

            anchor_cell, all_cells = resolve_target_cell(ws, cell_coord)
            is_merged = len(all_cells) > 1

            updated, highlighted = apply_cell_update(
                anchor_cell=anchor_cell,
                value=val,
                needs_review=needs_review,
                all_cells=all_cells,
            )

            if not updated:
                total_formulas_preserved += 1
                sheet_stats[sheet_name]["formulas"] += 1
                continue

            sheet_stats[sheet_name]["updated"] += 1
            if highlighted:
                sheet_stats[sheet_name]["highlighted"] += 1

            updated_records.append({
                "sheet_name": sheet_name,
                "cell_coordinate": cell_coord,
                "slot_type": slot_type,
                "value": val,
                "needs_review": needs_review,
                "highlighted": highlighted,
                "is_merged": is_merged,
            })

    output_excel_path.parent.mkdir(parents=True, exist_ok=True)
    saved_path = output_excel_path
    try:
        wb.save(output_excel_path)
        logger.info("Saved populated Excel to: %s", output_excel_path)
    except PermissionError as exc:
        timestamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
        saved_path = output_excel_path.parent / f"{output_excel_path.stem}_{timestamp}{output_excel_path.suffix}"
        logger.error(
            "Permission denied saving '%s' (file may be open in Excel). Saved to fallback '%s'. Error: %s",
            output_excel_path,
            saved_path,
            exc,
        )
        wb.save(saved_path)

    sheet_summaries: list[SheetSummary] = [
        {
            "sheet_name": s_name,
            "cells_updated": s_data["updated"],
            "cells_highlighted": s_data["highlighted"],
            "formulas_preserved": s_data["formulas"],
            "collisions_detected": s_data["collisions"],
        }
        for s_name, s_data in sheet_stats.items()
    ]

    unique_cell_keys = {(r["sheet_name"], r["cell_coordinate"]) for r in updated_records}

    manifest: PopulationManifest = {
        "document_id": doc_id,
        "workbook_id": wb_id,
        "source_file": str(original_excel_path),
        "output_excel_file": str(saved_path),
        "generated_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
        "total_requirements": len(items),
        "total_cells_updated": len(updated_records),
        "unique_cells_updated": len(unique_cell_keys),
        "total_cells_highlighted": sum(1 for r in updated_records if r["highlighted"]),
        "total_formulas_preserved": total_formulas_preserved,
        "total_collisions_detected": total_collisions,
        "sheet_summaries": sheet_summaries,
        "updated_cells": updated_records,
    }

    manifest_target = manifest_output_path or decisions_path.parent / "population_manifest.json"
    with open(manifest_target, "w", encoding="utf-8") as f:
        json.dump(manifest, f, indent=2)
    logger.info("Saved population manifest to: %s", manifest_target)

    return manifest


def populate_excel(
    template_path: Path | str,
    decisions_path: Path | str,
    output_excel_path: Path | str,
    manifest_path: Path | str | None = None,
) -> PopulationManifest:
    """Convenience wrapper for populate_workbook_in_place."""
    return populate_workbook_in_place(
        decisions_path=Path(decisions_path),
        original_excel_path=Path(template_path),
        output_excel_path=Path(output_excel_path),
        manifest_output_path=Path(manifest_path) if manifest_path else None,
    )


if __name__ == "__main__":
    import argparse
    import shutil

    parser = argparse.ArgumentParser(description="Step 5: In-Place Excel Workbook Populator")
    parser.add_argument(
        "--decisions",
        "-d",
        type=Path,
        default=Path(__file__).parent / "compliance_decisions.json",
        help="Path to Step 4 compliance decisions JSON",
    )
    parser.add_argument(
        "--analysis",
        "-a",
        type=Path,
        default=Path(__file__).parent / "excel_analysis.json",
        help="Path to Step 2 Excel analysis JSON",
    )
    parser.add_argument(
        "--template",
        "-t",
        type=Path,
        default=None,
        help="Optional explicit path to original Excel template",
    )
    parser.add_argument(
        "--output-excel",
        "-o",
        type=Path,
        default=None,
        help="Optional explicit path for populated output Excel file",
    )
    parser.add_argument(
        "--manifest",
        "-m",
        type=Path,
        default=Path(__file__).parent / "population_manifest.json",
        help="Path to output population manifest JSON",
    )
    args = parser.parse_args()

    base_dir = Path(__file__).parent
    decisions_file = args.decisions
    analysis_file = args.analysis

    orig_path: Path | None = args.template
    if orig_path is None and analysis_file.exists():
        with open(analysis_file, "r", encoding="utf-8") as af:
            af_data: dict[str, object] = json.load(af)
            src = af_data.get("source_file")
            if src:
                candidate = Path(str(src))
                if candidate.exists():
                    orig_path = candidate
                elif (candidate.parent / "input" / candidate.name).exists():
                    orig_path = candidate.parent / "input" / candidate.name
                elif (Path("data/uploads") / candidate.name).exists():
                    orig_path = Path("data/uploads") / candidate.name
                elif (Path("data") / candidate.name).exists():
                    orig_path = Path("data") / candidate.name

    if orig_path is None or not orig_path.exists():
        # Fallback: look for any .xlsx in data/uploads or data
        input_candidates = list(Path("data/uploads").glob("*.xlsx")) or list(Path("data").glob("*.xlsx"))
        if input_candidates:
            orig_path = input_candidates[0]
        else:
            raise FileNotFoundError(f"Source template not found in {analysis_file} or data/uploads. Please supply an original template.")

    # Determine output Excel path
    if args.output_excel:
        out_excel = args.output_excel
    elif orig_path.parent.name == "input" or orig_path.parent.name == "uploads":
        out_dir = Path("data/generated")
        out_dir.mkdir(parents=True, exist_ok=True)
        out_excel = out_dir / f"populated_{orig_path.name}"
    elif Path("data/generated").exists():
        out_excel = Path("data/generated") / f"populated_{orig_path.name}"
    else:
        out_excel = base_dir / f"populated_{orig_path.name}"

    manifest_file = args.manifest

    manifest_res = populate_workbook_in_place(
        decisions_path=decisions_file,
        original_excel_path=orig_path,
        output_excel_path=out_excel,
        manifest_output_path=manifest_file,
    )

    # Mirror copy to pipeline_lab/ if saved in data/generated/
    pipeline_copy = base_dir / f"populated_{orig_path.name}"
    if out_excel.resolve() != pipeline_copy.resolve():
        try:
            shutil.copy2(str(out_excel), str(pipeline_copy))
        except Exception:
            pass

    print(f"Step 5 Finished: Updated {manifest_res['total_cells_updated']} cells.")
    print(f"  -> Saved output Excel:    {out_excel.resolve()}")
    print(f"  -> Saved audit manifest:  {manifest_file.resolve()}")

