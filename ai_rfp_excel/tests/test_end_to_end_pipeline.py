from pathlib import Path

import openpyxl
import pytest

from ai_rfp_excel.app.ai.mock_provider import MockLLMProvider
from ai_rfp_excel.app.excel.analyzer import ExcelAnalyzer
from ai_rfp_excel.app.excel.validator import ExcelValidator
from ai_rfp_excel.app.excel.writer import ExcelWriter
from ai_rfp_excel.app.ingestion.pdf.router import PageRouter
from ai_rfp_excel.app.ingestion.pdf.text_extractor import extract_text_from_page
from ai_rfp_excel.app.matching.engine import ComplianceEngine
from ai_rfp_excel.app.matching.models import (
    ComplianceState,
    FactItem,
)


@pytest.mark.asyncio
async def test_complete_end_to_end_rfp_automation_pipeline(
    sample_pdf_text_only: Path,
    sample_excel_standard: Path,
    tmp_path: Path,
) -> None:
    """End-to-end integration test: PDF Ingestion -> Excel Analysis -> 5-Layer Matching -> Excel Population -> Validation."""

    # 1. Phase 1: PDF Ingestion & Specification Extraction
    router = PageRouter()
    page_types = router.detect_page_type(str(sample_pdf_text_only), 0)
    assert len(page_types) > 0

    page_text = extract_text_from_page(str(sample_pdf_text_only), 0)
    assert len(page_text.text) > 0

    extracted_facts: list[FactItem] = []
    # Create fact items from extracted lines/segments
    import re
    raw_lines = page_text.text.split("\n")
    all_segments: list[str] = []
    for line in raw_lines:
        segs = re.split(r"(?=(?:Model|Processors|Memory|Storage|Power|Network):)", line)
        all_segments.extend([s.strip() for s in segs if s.strip()])

    for seg in all_segments:
        if ":" in seg:
            parts = seg.split(":", 1)
            field = parts[0].strip()
            val = parts[1].strip()
            if field and val:
                extracted_facts.append(
                    FactItem(
                        field_name=field,
                        value=val,
                        source_page=1,
                        confidence=0.95,
                    )
                )

    assert len(extracted_facts) >= 4


    # 2. Phase 2: Excel Template Structure & Requirement Detection
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(
        file_path=str(sample_excel_standard),
        workbook_id="e2e-wb-001",
        version=1,
    )
    assert analysis.total_sheets == 1
    assert analysis.total_requirements == 5

    # 3. Phase 3: 5-Layer Compliance Resolution
    mock_llm = MockLLMProvider()
    engine = ComplianceEngine(llm_provider=mock_llm)

    decisions = []
    for sheet in analysis.sheets:
        for req in sheet.requirements:
            decision = await engine.evaluate_requirement(
                requirement_text=req.requirement_text,
                facts=extracted_facts,
                requirement_id=req.requirement_id,
            )
            decisions.append(decision)

    assert len(decisions) == 5

    # Check specific decisions
    dec_xeon = next(d for d in decisions if "Xeon" in d.requirement_text)
    assert dec_xeon.state == ComplianceState.COMPLIANT
    assert dec_xeon.resolving_layer in ("exact_match", "rule_based")

    dec_ram = next(d for d in decisions if "memory" in d.requirement_text.lower())
    assert dec_ram.state == ComplianceState.COMPLIANT
    assert dec_ram.resolving_layer in ("rule_based", "unit_conversion")

    # 4. Phase 4: Excel Population & Executive Summary Tab
    writer = ExcelWriter(output_dir=tmp_path / "e2e_output")
    result = writer.populate_workbook(
        template_path=sample_excel_standard,
        analysis=analysis,
        decisions=decisions,
    )

    assert Path(result.output_file_path).exists()
    assert result.summary_sheet_created is True
    assert result.total_populated_cells > 0

    # 5. Phase 5: Automated Quality Validation
    validator = ExcelValidator()
    val_report = validator.validate(
        populated_path=result.output_file_path,
        original_analysis=analysis,
        decisions=decisions,
    )

    assert val_report.is_valid is True
    assert val_report.passed_checks >= 5

    # 6. Verify populated workbook integrity
    pop_wb = openpyxl.load_workbook(result.output_file_path, data_only=False)
    assert pop_wb.sheetnames[0] == "Compliance Summary"
    assert "Technical Compliance" in pop_wb.sheetnames

    ws_data = pop_wb["Technical Compliance"]
    assert ws_data["C2"].value == "Compliant"  # Compliance cell
    assert ws_data["B7"].value == "=COUNT(A2:A6)"  # Formula preserved!


@pytest.mark.asyncio
async def test_end_to_end_renamed_columns_template(
    sample_excel_renamed_cols: Path,
    sample_facts_corpus: list[FactItem],
    tmp_path: Path,
) -> None:
    """Test end-to-end pipeline robustness against renamed and rearranged columns."""
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(
        file_path=str(sample_excel_renamed_cols),
        workbook_id="e2e-renamed-002",
        version=1,
    )
    assert analysis.total_requirements >= 1

    engine = ComplianceEngine(llm_provider=MockLLMProvider())
    decisions = []
    for sheet in analysis.sheets:
        for req in sheet.requirements:
            dec = await engine.evaluate_requirement(
                requirement_text=req.requirement_text,
                facts=sample_facts_corpus,
                requirement_id=req.requirement_id,
            )
            decisions.append(dec)

    writer = ExcelWriter(output_dir=tmp_path / "renamed_output")
    result = writer.populate_workbook(
        template_path=sample_excel_renamed_cols,
        analysis=analysis,
        decisions=decisions,
    )
    assert result.validation_report.is_valid is True
