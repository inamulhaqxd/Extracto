import os
import tempfile
from pathlib import Path
from uuid import uuid4

import openpyxl
import pytest
import pytest_asyncio
from sqlalchemy.ext.asyncio import AsyncSession

from app.database.models import User, Document, Workbook
from app.api.auth import get_password_hash
from app.excel.analyzer import ExcelAnalyzer
from app.matching.compliance import ComplianceEngine
from app.matching.rules import ComplianceStatus
from app.ai.mock_llm import MockLLMProvider


@pytest.fixture
def sample_pdf():
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Minimum 32 cores per controller\n64 GB RAM\n15TB NVMe Storage")

    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f:
        doc.save(f.name)
        yield f.name

    doc.close()
    os.unlink(f.name)


@pytest.fixture
def sample_excel():
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Requirements"
    ws["A1"] = "Requirement"
    ws["B1"] = "DWP"
    ws["C1"] = "Premier"
    ws["A2"] = "Minimum 32 cores per controller"
    ws["A3"] = "Minimum 64 GB RAM"
    ws["A4"] = "Minimum 15TB storage capacity"

    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as f:
        wb.save(f.name)
        yield f.name

    os.unlink(f.name)


@pytest.mark.asyncio
async def test_excel_analysis(sample_excel):
    analyzer = ExcelAnalyzer()
    result = analyzer.analyze(sample_excel)

    assert result["sheet_count"] == 1
    assert result["total_requirements"] > 0

    analyzer.close()


def test_compliance_engine():
    llm = MockLLMProvider()
    llm.set_default_response('{"status": "COMPLIANT", "confidence": 0.95, "evidence": [{"text": "32 cores found", "page": 1, "source_type": "text"}], "reasoning": "Match found"}')

    engine = ComplianceEngine(llm, "qwen3:4b")

    assert engine is not None
    assert engine.model == "qwen3:4b"


@pytest.mark.asyncio
async def test_mock_compliance_resolution():
    from app.matching.rules import ExactMatcher

    matcher = ExactMatcher()
    decision = matcher.match("Minimum 32 cores", "Minimum 32 cores")

    assert decision is not None
    assert decision.status == ComplianceStatus.COMPLIANT


@pytest.mark.asyncio
async def test_full_workflow(sample_excel):
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze(sample_excel)

    assert "requirements" in analysis
    assert len(analysis["requirements"]) > 0

    analyzer.close()
