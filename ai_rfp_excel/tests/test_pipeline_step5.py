"""
Unit tests for Step 5: "The In-Place Excel Populator & Styler".
Tests:
- Rule 9 In-Place Workbook & Style Preservation (fonts, borders, sheets).
- Formula Guard: Existing Excel formulas are NEVER overwritten.
- Rule 7 Visual Audit Highlighting: Soft yellow fill (FFF2CC) applied when needs_review=True.
- Valid coordinate enforcement & graceful handling of missing sheets.
- Population manifest generation and count accuracy (Rule 8).
- End-to-end population with real artifacts.
- Edge Case: Merged cell write protection and uniform highlight.
- Edge Case: Absolute coordinates ($B$4, $C$5) and lowercase (b4).
- Edge Case: Formula injection prevention (= 10 Gbps -> string data_type).
- Edge Case: None / empty value normalization to NOT_SPECIFIED with yellow fill.
- Edge Case: Case-insensitive and whitespace sheet resolution.
- Edge Case: Multiline text wrapping preservation.

Strict typing only — zero Any (Rule 4). 100% offline (Rule 10).
"""

from __future__ import annotations

import json
import sys
from pathlib import Path
from typing import cast

import openpyxl
import pytest
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from ai_rfp_excel.app.pipeline.step5_excel_populator import (
    YELLOW_HIGHLIGHT_HEX,
    apply_cell_update,
    find_sheet_case_insensitive,
    is_formula_cell,
    normalize_and_validate_coordinate,
    populate_workbook_in_place,
    resolve_target_cell,
)


def test_is_formula_cell() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws["A1"] = "Plain Text"
    ws["A2"] = "=SUM(B1:B10)"
    ws["A3"] = 12345

    assert not is_formula_cell(ws["A1"])
    assert is_formula_cell(ws["A2"])
    assert not is_formula_cell(ws["A3"])


def test_formula_protection_guard() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws["C5"] = "=A5*B5"

    updated, highlighted = apply_cell_update(ws["C5"], "OVERWRITE_ATTEMPT", needs_review=True)
    assert not updated
    assert not highlighted
    assert ws["C5"].value == "=A5*B5"  # Formula was strictly protected


def test_cell_update_value_and_yellow_highlight() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["B4"]
    cell.value = "Original"

    updated, highlighted = apply_cell_update(cell, "NOT_SPECIFIED", needs_review=True)
    assert updated
    assert highlighted
    assert cell.value == "NOT_SPECIFIED"

    # Verify yellow fill
    fill = cast(PatternFill, cell.fill)
    assert fill.fill_type == "solid"
    color_code = fill.start_color.rgb if hasattr(fill.start_color, "rgb") else str(fill.start_color)
    assert YELLOW_HIGHLIGHT_HEX in str(color_code)


def test_cell_update_without_highlight_when_verified() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["B4"]

    updated, highlighted = apply_cell_update(cell, "150", needs_review=False)
    assert updated
    assert not highlighted
    assert cell.value == "150"
    assert cell.fill.fill_type is None


def test_preserve_existing_cell_font_and_border() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["B10"]
    cell.font = Font(name="Arial", size=14, bold=True, color="FF0000")
    thin_side = Side(border_style="thin", color="000000")
    cell.border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    updated, highlighted = apply_cell_update(cell, "Verified Value", needs_review=False)
    assert updated
    assert not highlighted
    assert cell.value == "Verified Value"
    assert cell.font.name == "Arial"
    assert cell.font.size == 14
    assert cell.font.bold is True
    assert cell.border.left.border_style == "thin"


def test_merged_cell_resolution_and_styling() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.merge_cells("B2:C3")

    # Target the non-top-left cell of the merged range
    anchor, all_cells = resolve_target_cell(ws, "C3")
    assert anchor.coordinate == "B2"
    assert len(all_cells) == 4

    updated, highlighted = apply_cell_update(anchor, "Merged Answer", needs_review=True, all_cells=all_cells)
    assert updated
    assert highlighted
    assert anchor.value == "Merged Answer"

    # Check that all cells in the merged range receive yellow highlight
    for c in all_cells:
        fill = cast(PatternFill, c.fill)
        assert fill.fill_type == "solid"
        color_code = fill.start_color.rgb if hasattr(fill.start_color, "rgb") else str(fill.start_color)
        assert YELLOW_HIGHLIGHT_HEX in str(color_code)


def test_coordinate_normalization_and_validation() -> None:
    assert normalize_and_validate_coordinate("b4") == "B4"
    assert normalize_and_validate_coordinate("$B$4") == "B4"
    assert normalize_and_validate_coordinate(" $AA$100 ") == "AA100"
    assert normalize_and_validate_coordinate("INVALID") is None
    assert normalize_and_validate_coordinate("A999999999") is None  # Out of Excel bounds
    assert normalize_and_validate_coordinate("") is None


def test_formula_injection_prevention() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["A1"]

    # Specification starting with = or -
    apply_cell_update(cell, "= 10 Gbps minimum bandwidth", needs_review=False)
    assert cell.value == "= 10 Gbps minimum bandwidth"
    assert cell.data_type == "s"  # Protected as string, not parsed as Excel formula


def test_none_value_enforces_not_specified_and_yellow_fill() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["A1"]

    # Missing value normalized to NOT_SPECIFIED
    updated, highlighted = apply_cell_update(cell, "", needs_review=False)
    assert updated
    assert highlighted
    assert cell.value == "NOT_SPECIFIED"
    assert cell.fill.fill_type == "solid"


def test_find_sheet_case_insensitive() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.title = "Technical Specs"

    found_ws, canonical_name = find_sheet_case_insensitive(wb, "technical specs")
    assert canonical_name == "Technical Specs"
    assert found_ws.title == "Technical Specs"

    _found_ws2, canonical_name2 = find_sheet_case_insensitive(wb, "  Technical Specs  ")
    assert canonical_name2 == "Technical Specs"


def test_multiline_text_wrapping() -> None:
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    cell = ws["A1"]

    multiline_text = "Line 1: Requirement met.\nLine 2: Section 4.2 cited."
    apply_cell_update(cell, multiline_text, needs_review=False)
    assert cell.alignment.wrap_text is True


def test_populate_workbook_in_place_end_to_end(tmp_path: Path) -> None:
    # 1. Create a template workbook with a merged cell
    template_wb = openpyxl.Workbook()
    ws: Worksheet = template_wb.active
    ws.title = "Extraction"
    ws["A1"] = "Header Info"
    ws["A3"] = "Question"
    ws["B3"] = "AI Answer"
    ws["C3"] = "Remarks"
    ws["D3"] = "Calculated Total"

    ws["A4"] = "What is the warranty period?"
    ws["B4"] = None
    ws["C4"] = None

    ws["A5"] = "Is ISO certification provided?"
    ws["B5"] = None
    ws["C5"] = None

    ws["D4"] = "=SUM(10, 20)"  # Formula cell

    # Merge B5:C5 for testing merged cell in end-to-end
    ws.merge_cells("B5:C5")

    template_file = tmp_path / "test_template.xlsx"
    template_wb.save(template_file)

    # 2. Create mock decisions json with edge cases ($B$4 coordinate, null value)
    decisions_data: dict[str, object] = {
        "document_id": "TEST-DOC-5",
        "workbook_id": "TEST-WB-5",
        "total_requirements": 2,
        "items": [
            {
                "requirement_id": "REQ-01",
                "sheet_name": "extraction",  # Lowercase sheet name test
                "row_number": 4,
                "slot_assignments": {
                    "answer": {
                        "slot_type": "answer",
                        "cell_coordinate": "$B$4",  # Absolute coordinate test
                        "value": "3 Years",
                        "needs_review": False,
                    },
                    "remarks": {
                        "slot_type": "remarks",
                        "cell_coordinate": "C4",
                        "value": "Standard warranty applies. (Page 2)",
                        "needs_review": False,
                    },
                    "formula_col": {
                        "slot_type": "calculated",
                        "cell_coordinate": "D4",
                        "value": "OVERWRITE_ATTEMPT",
                        "needs_review": False,
                    },
                },
            },
            {
                "requirement_id": "REQ-02",
                "sheet_name": "Extraction",
                "row_number": 5,
                "slot_assignments": {
                    "answer": {
                        "slot_type": "answer",
                        "cell_coordinate": "C5",  # Target inside merged B5:C5
                        "value": None,  # None value test
                        "needs_review": False,
                    },
                },
            },
        ],
    }

    decisions_file = tmp_path / "compliance_decisions.json"
    with open(decisions_file, "w", encoding="utf-8") as f:
        json.dump(decisions_data, f)

    output_excel = tmp_path / "populated_output.xlsx"
    manifest_file = tmp_path / "population_manifest.json"

    # 3. Execute Step 5
    manifest = populate_workbook_in_place(
        decisions_path=decisions_file,
        original_excel_path=template_file,
        output_excel_path=output_excel,
        manifest_output_path=manifest_file,
    )

    # 4. Verify manifest
    assert manifest["document_id"] == "TEST-DOC-5"
    assert manifest["total_requirements"] == 2
    assert manifest["total_cells_updated"] == 3  # B4, C4, and B5 (via C5)
    assert manifest["total_cells_highlighted"] == 1  # B5 (due to None -> NOT_SPECIFIED)
    assert manifest["total_formulas_preserved"] == 1  # D4
    assert manifest_file.exists()

    # 5. Reload populated workbook and verify contents
    loaded_wb = openpyxl.load_workbook(output_excel, data_only=False)
    loaded_ws: Worksheet = loaded_wb["Extraction"]

    # REQ-01 verified
    assert loaded_ws["B4"].value == "3 Years"
    assert loaded_ws["C4"].value == "Standard warranty applies. (Page 2)"
    assert loaded_ws["B4"].fill.fill_type is None
    assert loaded_ws["D4"].value == "=SUM(10, 20)"  # Formula preserved

    # REQ-02 merged B5:C5 has NOT_SPECIFIED on anchor B5 and yellow highlight
    assert loaded_ws["B5"].value == "NOT_SPECIFIED"
    fill_b5 = cast(PatternFill, loaded_ws["B5"].fill)
    assert fill_b5.fill_type == "solid"
    color_b5 = fill_b5.start_color.rgb if hasattr(fill_b5.start_color, "rgb") else str(fill_b5.start_color)
    assert YELLOW_HIGHLIGHT_HEX in str(color_b5)


def test_missing_files_raise_filenotfound(tmp_path: Path) -> None:
    non_existent = tmp_path / "does_not_exist.json"
    with pytest.raises(FileNotFoundError):
        populate_workbook_in_place(
            decisions_path=non_existent,
            original_excel_path=tmp_path / "none.xlsx",
            output_excel_path=tmp_path / "out.xlsx",
        )

