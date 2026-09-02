from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font, PatternFill

from ai_rfp_excel.app.ai.mock_provider import MockLLMProvider
from ai_rfp_excel.app.excel.analyzer import ExcelAnalyzer
from ai_rfp_excel.app.excel.writer import ExcelWriter
from ai_rfp_excel.app.matching.engine import ComplianceEngine
from ai_rfp_excel.app.matching.models import (
    ComplianceDecision,
    ComplianceState,
    EvidenceItem,
    FactItem,
)


def test_edge_case_custom_offered_spec_and_empty_slots(tmp_path: Path) -> None:
    """Edge Case: Complex sheet with 'Offered Specification', 'Make / Model', 'Compliance', 'Remarks', and empty slots."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Custom Matrix"

    # Multi-slot header with non-standard names
    headers = [
        "S.No",
        "Technical Requirement Description",
        "Offered Specification / Parameter",
        "Make & Model Quoted",
        "Compliance (Complied/Not Complied)",
        "Deviations & Remarks",
    ]
    for col_idx, h in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=h)

    # Row 2: Empty slots
    ws.cell(row=2, column=1, value="1")
    ws.cell(row=2, column=2, value="Intel Xeon Gold 6430 CPU")
    ws.cell(row=2, column=3, value="")
    ws.cell(row=2, column=4, value="")
    ws.cell(row=2, column=5, value="")
    ws.cell(row=2, column=6, value="")

    # Row 3: Partially pre-filled slot (e.g. user already filled vendor model)
    ws.cell(row=3, column=1, value="2")
    ws.cell(row=3, column=2, value="512GB DDR5 Registered RAM")
    ws.cell(row=3, column=3, value="")
    ws.cell(row=3, column=4, value="Samsung M321R8GA0BB0")
    ws.cell(row=3, column=5, value="")
    ws.cell(row=3, column=6, value="")

    file_path = tmp_path / "custom_matrix.xlsx"
    wb.save(file_path)

    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(file_path))

    assert analysis.total_sheets == 1
    assert analysis.total_requirements == 2

    req1 = analysis.sheets[0].requirements[0]
    req2 = analysis.sheets[0].requirements[1]

    # Check that offered spec columns and vendor columns are recognized
    assert len(req1.offered_spec_cells) >= 1
    assert len(req1.compliance_cells) >= 1
    assert len(req1.remarks_cells) >= 1

    # Row 2 has 4 empty slots (C2, D2, E2, F2)
    assert "C2" in req1.empty_slots
    assert "D2" in req1.empty_slots
    assert "E2" in req1.empty_slots
    assert "F2" in req1.empty_slots

    # Row 3 has D3 pre-filled, so D3 is not in empty_slots
    assert "C3" in req2.empty_slots
    assert "D3" not in req2.empty_slots
    assert "E3" in req2.empty_slots
    assert "F3" in req2.empty_slots


def test_edge_case_merged_cells_and_blank_rows(tmp_path: Path) -> None:
    """Edge Case: Worksheets with merged cells across titles and sporadic blank rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Merged Sheet"

    # Row 1: Merged Title Block
    ws.merge_cells("A1:D1")
    ws["A1"] = "PROJECT TENDER SPECIFICATION MATRIX"

    # Row 2: Empty spacer row
    # Row 3: Table Headers
    ws.cell(row=3, column=1, value="Item No")
    ws.cell(row=3, column=2, value="Specification Clause")
    ws.cell(row=3, column=3, value="Offered Value")
    ws.cell(row=3, column=4, value="Compliance Status")

    # Row 4: Section Header
    ws.cell(row=4, column=1, value="Section 1.0 - Compute Nodes")

    # Row 5: Requirement
    ws.cell(row=5, column=1, value="1.1")
    ws.cell(row=5, column=2, value="Dual Socket Intel Xeon CPU")
    ws.cell(row=5, column=3, value="")
    ws.cell(row=5, column=4, value="")

    # Row 6: Blank row
    # Row 7: Requirement
    ws.cell(row=7, column=1, value="1.2")
    ws.cell(row=7, column=2, value="150TB All-Flash NVMe SSDs")
    ws.cell(row=7, column=3, value="")
    ws.cell(row=7, column=4, value="")

    file_path = tmp_path / "merged_sheet.xlsx"
    wb.save(file_path)

    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(file_path))

    assert analysis.total_requirements == 2
    req_texts = [r.requirement_text for r in analysis.sheets[0].requirements]
    assert "Dual Socket Intel Xeon CPU" in req_texts
    assert "150TB All-Flash NVMe SSDs" in req_texts


def test_edge_case_multi_sheet_diverse_layouts(tmp_path: Path) -> None:
    """Edge Case: Multi-sheet workbook with different column structures per sheet."""
    wb = openpyxl.Workbook()

    # Sheet 1: Servers
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Servers"
    ws1.cell(row=1, column=1, value="Req #")
    ws1.cell(row=1, column=2, value="Description")
    ws1.cell(row=1, column=3, value="Offered Spec")
    ws1.cell(row=1, column=4, value="Status")

    ws1.cell(row=2, column=1, value="S1")
    ws1.cell(row=2, column=2, value="2x Intel Xeon Gold 6430")
    ws1.cell(row=2, column=3, value="")
    ws1.cell(row=2, column=4, value="")

    # Sheet 2: Networking
    ws2 = wb.create_sheet(title="Networking")
    ws2.cell(row=1, column=1, value="Item")
    ws2.cell(row=1, column=2, value="Network Criteria")
    ws2.cell(row=1, column=3, value="Compliance")
    ws2.cell(row=1, column=4, value="Remarks")

    ws2.cell(row=2, column=1, value="N1")
    ws2.cell(row=2, column=2, value="4x 25GbE SFP28 ports")
    ws2.cell(row=2, column=3, value="")
    ws2.cell(row=2, column=4, value="")

    file_path = tmp_path / "multi_sheet.xlsx"
    wb.save(file_path)

    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(file_path))

    assert analysis.total_sheets == 2
    assert analysis.total_requirements == 2
    assert analysis.sheets[0].sheet_name == "Servers"
    assert analysis.sheets[1].sheet_name == "Networking"


def test_edge_case_writer_populates_offered_specs_and_preserves_formulas(tmp_path: Path) -> None:
    """Edge Case: Verify writer populates offered specs, compliance styling, and preserves formulas & formatting."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Compliance Matrix"

    # Header with styling
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")

    headers = ["ID", "Technical Requirement", "Offered Specification", "Compliance Status", "Remarks"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font

    # Rows with formula
    ws.cell(row=2, column=1, value="REQ-01")
    ws.cell(row=2, column=2, value="Intel Xeon Gold 6430")
    ws.cell(row=2, column=3, value="")
    ws.cell(row=2, column=4, value="")
    ws.cell(row=2, column=5, value="")

    ws.cell(row=3, column=1, value="REQ-02")
    ws.cell(row=3, column=2, value="512GB DDR5 RAM")
    ws.cell(row=3, column=3, value="")
    ws.cell(row=3, column=4, value="")
    ws.cell(row=3, column=5, value="")

    # Formula row at bottom
    ws.cell(row=4, column=1, value="Count")
    ws.cell(row=4, column=2, value="=COUNTA(B2:B3)")

    template_path = tmp_path / "formula_template.xlsx"
    wb.save(template_path)

    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(template_path))

    decisions = [
        ComplianceDecision(
            requirement_id="REQ-S01-001",
            requirement_text="Intel Xeon Gold 6430",
            state=ComplianceState.COMPLIANT,
            confidence=0.98,
            reasoning="Exact processor match found in datasheet.",
            matched_value="Dual Intel Xeon Gold 6430 32-Core 2.1GHz",
            resolving_layer="exact_match",
            evidence=[
                EvidenceItem(
                    value="Intel Xeon Gold 6430",
                    confidence=0.98,
                    extraction_method="exact",
                    citation="Page 2",
                    reasoning="Processor spec verified",
                )
            ],
        ),
        ComplianceDecision(
            requirement_id="REQ-S01-002",
            requirement_text="512GB DDR5 RAM",
            state=ComplianceState.COMPLIANT,
            confidence=0.95,
            reasoning="Memory exceeds minimum threshold.",
            matched_value="512GB DDR5-4800 ECC Registered",
            resolving_layer="rule_based",
            evidence=[
                EvidenceItem(
                    value="512GB DDR5",
                    confidence=0.95,
                    extraction_method="rule_based",
                    citation="Page 3",
                    reasoning="RAM spec verified",
                )
            ],
        ),
    ]

    writer = ExcelWriter(output_dir=tmp_path / "edge_output")
    result = writer.populate_workbook(
        template_path=template_path,
        analysis=analysis,
        decisions=decisions,
    )

    assert Path(result.output_file_path).exists()
    assert result.total_populated_cells > 0

    # Load output workbook to verify cell values & formula
    res_wb = openpyxl.load_workbook(result.output_file_path, data_only=False)
    res_ws = res_wb["Compliance Matrix"]

    # Check offered specification populated
    assert res_ws["C2"].value == "Dual Intel Xeon Gold 6430 32-Core 2.1GHz"
    assert res_ws["C3"].value == "512GB DDR5-4800 ECC Registered"

    # Check compliance status populated
    assert res_ws["D2"].value == "Compliant"
    assert res_ws["D3"].value == "Compliant"

    # Check remarks populated with citation
    assert "Page 2" in str(res_ws["E2"].value)
    assert "Page 3" in str(res_ws["E3"].value)

    # Check formula preserved
    assert res_ws["B4"].value == "=COUNTA(B2:B3)"

    res_wb.close()


def test_edge_case_special_characters_and_long_text(tmp_path: Path) -> None:
    """Edge Case: Handles special characters, quotes, and long strings safely."""
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Special Chars"

    ws.cell(row=1, column=1, value="Requirement")
    ws.cell(row=1, column=2, value="Offered Spec")
    ws.cell(row=1, column=3, value="Compliance")
    ws.cell(row=1, column=4, value="Remarks")

    special_req = "Redundant Power (2x 1600W, >=96% Eff., ~240V ±10% / 50-60Hz) & 'Hot-Swap' Support <10ms"
    ws.cell(row=2, column=1, value=special_req)
    ws.cell(row=2, column=2, value="")
    ws.cell(row=2, column=3, value="")
    ws.cell(row=2, column=4, value="")

    template_path = tmp_path / "special_template.xlsx"
    wb.save(template_path)

    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(template_path))

    long_reasoning = "Detailed justification: " + "Specification complies with all industrial standards. " * 15

    decisions = [
        ComplianceDecision(
            requirement_id="REQ-S01-001",
            requirement_text=special_req,
            state=ComplianceState.COMPLIANT,
            confidence=0.92,
            reasoning=long_reasoning,
            matched_value="2x 1600W Titanium 96% Eff. Hot-Swap PSU",
            resolving_layer="rule_based",
            evidence=[
                EvidenceItem(
                    value="2x 1600W Titanium",
                    confidence=0.92,
                    extraction_method="rules",
                    citation="Page 6, Table 4",
                    reasoning=long_reasoning,
                )
            ],
        )
    ]

    writer = ExcelWriter(output_dir=tmp_path / "special_output")
    result = writer.populate_workbook(template_path, analysis, decisions)

    assert Path(result.output_file_path).exists()
    res_wb = openpyxl.load_workbook(result.output_file_path, data_only=True)
    res_ws = res_wb["Special Chars"]

    assert res_ws["B2"].value == "2x 1600W Titanium 96% Eff. Hot-Swap PSU"
    assert res_ws["C2"].value == "Compliant"
    assert "Page 6, Table 4" in str(res_ws["D2"].value)
    res_wb.close()


@pytest.mark.asyncio
async def test_edge_case_zero_facts_hallucination_protection() -> None:
    """Edge Case: When zero facts are extracted from PDF, all requirements must resolve to NOT_FOUND with zero hallucination."""
    mock_llm = MockLLMProvider()
    engine = ComplianceEngine(llm_provider=mock_llm)

    empty_facts: list[FactItem] = []

    dec = await engine.evaluate_requirement(
        requirement_text="128GB DDR5 RAM",
        facts=empty_facts,
        requirement_id="REQ-EMPTY-001",
    )

    assert dec.state == ComplianceState.NOT_FOUND
    assert dec.confidence >= 0.90
    assert len(dec.evidence) == 0
    assert "No technical specifications or facts found" in dec.reasoning or "not found" in dec.reasoning.lower()


@pytest.mark.asyncio
async def test_edge_case_conflicting_facts_resolution() -> None:
    """Edge Case: Conflicting specifications on different pages produce AMBIGUOUS compliance state."""
    mock_llm = MockLLMProvider()
    engine = ComplianceEngine(llm_provider=mock_llm)

    conflicting_facts = [
        FactItem(
            field_name="System Memory",
            value="128GB RAM",
            source_page=2,
            confidence=0.90,
        ),
        FactItem(
            field_name="System Memory",
            value="256GB RAM",
            source_page=5,
            confidence=0.90,
        ),
    ]

    dec = await engine.evaluate_requirement(
        requirement_text="Provide System Memory specifications",
        facts=conflicting_facts,
        requirement_id="REQ-CONFLICT-001",
    )

    assert dec.state == ComplianceState.AMBIGUOUS
    assert len(dec.conflicting_evidence) >= 2
    assert "conflicting" in dec.reasoning.lower()
