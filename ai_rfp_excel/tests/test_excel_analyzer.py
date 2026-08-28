import tempfile
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
import pytest

from ai_rfp_excel.app.excel.analyzer import ExcelAnalyzer
from ai_rfp_excel.app.excel.models import ColumnType
from ai_rfp_excel.app.excel.utils import (
    calculate_workbook_hash,
    validate_excel_file,
)


def create_sample_workbook(file_path: str) -> None:
    wb = openpyxl.Workbook()

    # Sheet 1: Servers
    ws1 = wb.active
    assert ws1 is not None
    ws1.title = "Servers"

    # Header in Row 2
    ws1.cell(row=2, column=1, value="S.No")
    ws1.cell(row=2, column=2, value="Technical Specification")
    ws1.cell(row=2, column=3, value="Vendor A")
    ws1.cell(row=2, column=4, value="Compliance (Yes/No)")
    ws1.cell(row=2, column=5, value="Remarks")

    # Section 1 Header
    ws1.cell(row=3, column=1, value="1.0 Compute Infrastructure")
    ws1.cell(row=3, column=1).font = Font(bold=True)

    # Requirement 1
    ws1.cell(row=4, column=1, value="1.1")
    ws1.cell(row=4, column=2, value="Minimum 2x Intel Xeon Gold 6430 Processors")
    ws1.cell(row=4, column=3, value="2x Intel Xeon Gold 6430")
    ws1.cell(row=4, column=4, value="Yes")
    ws1.cell(row=4, column=5, value="Fully compliant")

    # Requirement 2
    ws1.cell(row=5, column=1, value="1.2")
    ws1.cell(row=5, column=2, value="At least 256GB DDR5 ECC Registered Memory")

    # Merged title cell
    ws1.merge_cells("A1:E1")
    ws1.cell(row=1, column=1, value="RFP BOM Specification")

    # Hidden Row & Hidden Column
    ws1.row_dimensions[6].hidden = True
    ws1.column_dimensions["F"].hidden = True

    # Sheet 2: Storage
    ws2 = wb.create_sheet(title="Storage")
    ws2.cell(row=1, column=1, value="Item Description")
    ws2.cell(row=1, column=2, value="Dell PowerStore")
    ws2.cell(row=1, column=3, value="HPE Alletra")
    ws2.cell(row=1, column=4, value="Status")

    ws2.cell(row=2, column=1, value="NVMe All-Flash Array with >= 100TB usable")
    ws2.cell(row=2, column=2, value="PowerStore 5000T")
    ws2.cell(row=2, column=3, value="Alletra 9000")
    ws2.cell(row=2, column=4, value="Complied")

    ws2.cell(row=3, column=1, value="Dual Active-Active Controllers")

    wb.save(file_path)
    wb.close()


def test_excel_analyzer_basic_structure() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "test_rfp.xlsx"
        create_sample_workbook(str(excel_path))

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path), workbook_id="wb-101", version=1)

        assert analysis.workbook_id == "wb-101"
        assert analysis.total_sheets == 2
        assert analysis.sheet_names == ["Servers", "Storage"]
        assert len(analysis.file_hash) == 64
        assert analysis.file_size > 0
        assert analysis.total_requirements == 4  # 2 in Servers, 2 in Storage


def test_excel_analyzer_sheet_metadata_and_hidden() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "test_rfp.xlsx"
        create_sample_workbook(str(excel_path))

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path))

        servers_sheet = analysis.sheets[0]
        assert servers_sheet.sheet_name == "Servers"
        assert servers_sheet.sheet_index == 0

        # Dimensions
        assert servers_sheet.dimensions["max_row"] >= 5
        assert servers_sheet.dimensions["max_column"] >= 5

        # Merged cells
        assert len(servers_sheet.merged_cells) >= 1
        merged_ranges = [m["range"] for m in servers_sheet.merged_cells]
        assert "A1:E1" in merged_ranges

        # Hidden row & column
        assert 6 in servers_sheet.hidden_rows
        assert "F" in servers_sheet.hidden_columns


def test_excel_analyzer_column_detection() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "test_rfp.xlsx"
        create_sample_workbook(str(excel_path))

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path))

        servers_sheet = analysis.sheets[0]
        cols_by_type = {c.column_type: c for c in servers_sheet.columns}

        assert ColumnType.INDEX in cols_by_type
        assert ColumnType.REQUIREMENT in cols_by_type
        assert ColumnType.VENDOR in cols_by_type
        assert ColumnType.COMPLIANCE in cols_by_type
        assert ColumnType.REMARKS in cols_by_type

        # Verify Storage sheet with multiple vendor columns
        storage_sheet = analysis.sheets[1]
        vendor_cols = [c for c in storage_sheet.columns if c.column_type == ColumnType.VENDOR]
        assert len(vendor_cols) == 2
        vendor_names = {c.header_name for c in vendor_cols}
        assert "Dell PowerStore" in vendor_names
        assert "HPE Alletra" in vendor_names


def test_excel_analyzer_sections_and_requirements() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "test_rfp.xlsx"
        create_sample_workbook(str(excel_path))

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path))

        servers_sheet = analysis.sheets[0]
        assert len(servers_sheet.sections) == 1
        assert "1.0 Compute Infrastructure" in servers_sheet.sections[0].name

        assert len(servers_sheet.requirements) == 2
        req1 = servers_sheet.requirements[0]
        assert req1.requirement_id == "REQ-S01-001"
        assert "Intel Xeon" in req1.requirement_text
        assert req1.source_cell == "B4"
        assert req1.section == "1.0 Compute Infrastructure"
        assert "Vendor A" in req1.vendor_cells
        assert req1.vendor_cells["Vendor A"] == "C4"
        assert req1.compliance_cells["Compliance (Yes/No)"] == "D4"
        assert req1.remarks_cells["Remarks"] == "E4"

        req2 = servers_sheet.requirements[1]
        assert req2.requirement_id == "REQ-S01-002"
        assert "256GB DDR5" in req2.requirement_text
        assert req2.source_cell == "B5"


def test_excel_analyzer_renamed_columns_and_rearranged_order() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "dynamic_rfp.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "CustomLayout"

        # Headers rearranged: Remarks in Col A, Specs in Col B, Supplier in Col C, Status in Col D
        ws.cell(row=1, column=1, value="Evidence / Notes")
        ws.cell(row=1, column=2, value="Scope & Criteria")
        ws.cell(row=1, column=3, value="Premier Solution")
        ws.cell(row=1, column=4, value="Conformity")

        ws.cell(row=2, column=1, value="Refer doc page 10")
        ws.cell(row=2, column=2, value="Must support 100GbE QSFP28 Uplinks")
        ws.cell(row=2, column=3, value="SN2410 Switch")
        ws.cell(row=2, column=4, value="Yes")

        wb.save(str(excel_path))
        wb.close()

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path))

        sheet = analysis.sheets[0]
        reqs = sheet.requirements
        assert len(reqs) == 1
        assert reqs[0].source_cell == "B2"
        assert "100GbE QSFP28" in reqs[0].requirement_text
        assert reqs[0].vendor_cells["Premier Solution"] == "C2"
        assert reqs[0].compliance_cells["Conformity"] == "D2"
        assert reqs[0].remarks_cells["Evidence / Notes"] == "A2"


def test_validate_excel_file() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        valid_path = Path(tmp_dir) / "test.xlsx"
        create_sample_workbook(str(valid_path))

        # Valid file
        valid, err = validate_excel_file(str(valid_path), max_size_mb=50)
        assert valid is True
        assert err is None

        # Invalid extension
        invalid_ext_path = Path(tmp_dir) / "test.pdf"
        invalid_ext_path.write_bytes(b"some content")
        valid, err = validate_excel_file(str(invalid_ext_path))
        assert valid is False
        assert "Invalid file extension" in (err or "")

        # Non-existent file
        valid, err = validate_excel_file(str(Path(tmp_dir) / "missing.xlsx"))
        assert valid is False
        assert "does not exist" in (err or "")

        # Oversized file simulation (max size 0MB threshold)
        valid, err = validate_excel_file(str(valid_path), max_size_mb=0)
        assert valid is False
        assert "exceeds maximum allowed limit" in (err or "")


def test_calculate_workbook_hash() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        p1 = Path(tmp_dir) / "wb1.xlsx"
        p2 = Path(tmp_dir) / "wb1_copy.xlsx"
        create_sample_workbook(str(p1))
        p2.write_bytes(p1.read_bytes())

        hash1 = calculate_workbook_hash(str(p1))
        hash2 = calculate_workbook_hash(str(p2))
        assert len(hash1) == 64
        assert hash1 == hash2



def test_excel_analyzer_added_and_removed_sheets() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        base_path = Path(tmp_dir) / "base.xlsx"
        modified_path = Path(tmp_dir) / "modified.xlsx"

        # Create base with 2 sheets
        create_sample_workbook(str(base_path))

        # Create modified with 3 sheets (added "Networking", removed "Storage")
        wb = openpyxl.Workbook()
        ws1 = wb.active
        assert ws1 is not None
        ws1.title = "Servers"
        ws1.cell(row=1, column=1, value="Requirement")
        ws1.cell(row=2, column=1, value="2x Intel Xeon Gold")

        ws2 = wb.create_sheet(title="Networking")
        ws2.cell(row=1, column=1, value="Specification")
        ws2.cell(row=1, column=2, value="Vendor Cisco")
        ws2.cell(row=2, column=1, value="48-Port PoE+ Layer 3 Switch")
        ws2.cell(row=2, column=2, value="Catalyst 9300")

        ws3 = wb.create_sheet(title="Security")
        ws3.cell(row=1, column=1, value="Description")
        ws3.cell(row=2, column=1, value="Next-Gen Firewall with 10Gbps IPS")

        wb.save(str(modified_path))
        wb.close()

        analyzer = ExcelAnalyzer()
        base_analysis = analyzer.analyze_workbook(str(base_path))
        mod_analysis = analyzer.analyze_workbook(str(modified_path))

        assert base_analysis.total_sheets == 2
        assert "Storage" in base_analysis.sheet_names

        assert mod_analysis.total_sheets == 3
        assert "Storage" not in mod_analysis.sheet_names
        assert "Networking" in mod_analysis.sheet_names
        assert "Security" in mod_analysis.sheet_names
        assert mod_analysis.total_requirements == 3


def test_excel_analyzer_complex_multi_sections() -> None:
    with tempfile.TemporaryDirectory() as tmp_dir:
        excel_path = Path(tmp_dir) / "multi_section.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "ComplexSpecs"

        # Header row
        ws.cell(row=1, column=1, value="ID")
        ws.cell(row=1, column=2, value="Technical Requirement")
        ws.cell(row=1, column=3, value="OEM Model")
        ws.cell(row=1, column=4, value="Complied")

        # Section 1
        ws.cell(row=2, column=1, value="Section 1: High Availability")
        ws.cell(row=3, column=1, value="HA-01")
        ws.cell(row=3, column=2, value="Redundant hot-swappable power supplies")
        ws.cell(row=3, column=3, value="Dual 1600W Titanium")
        ws.cell(row=3, column=4, value="Yes")

        # Section 2
        ws.cell(row=4, column=1, value="Section 2: Management & Monitoring")
        ws.cell(row=5, column=1, value="MGT-01")
        ws.cell(row=5, column=2, value="Dedicated out-of-band IPMI 2.0 interface")
        ws.cell(row=5, column=3, value="iDRAC9 Enterprise")
        ws.cell(row=5, column=4, value="Yes")

        ws.cell(row=6, column=1, value="MGT-02")
        ws.cell(row=6, column=2, value="SNMPv3 and RESTful API support")
        ws.cell(row=6, column=3, value="Supported")
        ws.cell(row=6, column=4, value="Yes")

        wb.save(str(excel_path))
        wb.close()

        analyzer = ExcelAnalyzer()
        analysis = analyzer.analyze_workbook(str(excel_path))
        sheet = analysis.sheets[0]

        assert len(sheet.sections) == 2
        assert sheet.sections[0].name == "Section 1: High Availability"
        assert len(sheet.sections[0].requirements) == 1
        assert sheet.sections[1].name == "Section 2: Management & Monitoring"
        assert len(sheet.sections[1].requirements) == 2
        assert len(sheet.requirements) == 3

