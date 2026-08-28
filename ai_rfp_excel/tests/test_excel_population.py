import openpyxl
from openpyxl.styles import Border, Font, PatternFill, Side
from pathlib import Path
import pytest

from ai_rfp_excel.app.excel.analyzer import ExcelAnalyzer
from ai_rfp_excel.app.excel.validator import ExcelValidator
from ai_rfp_excel.app.excel.writer import ExcelWriter
from ai_rfp_excel.app.matching.models import (
    ComplianceDecision,
    ComplianceState,
    EvidenceItem,
)


@pytest.fixture
def sample_rfp_workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()

    # Sheet 1: Server Requirements
    ws1 = wb.active
    ws1.title = "Server Specs"

    # Header and styles
    bold_font = Font(name="Arial", size=12, bold=True)
    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    fill_header = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    headers = ["Item #", "Specification / Requirement", "Compliance (Yes/No)", "Remarks / Proof"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws1.cell(row=1, column=col_idx, value=h)
        cell.font = bold_font
        cell.fill = fill_header
        cell.border = border

    # Requirement rows
    requirements = [
        (1, "Minimum 256GB DDR5 ECC RAM required"),
        (2, "Dual redundant 1600W Titanium power supplies"),
        (3, "At least 100TB All-Flash NVMe usable storage"),
        (4, "Liquid immersion cooling system"),
    ]

    for idx, (num, req_text) in enumerate(requirements, start=2):
        ws1.cell(row=idx, column=1, value=num).font = Font(name="Arial", size=10)
        ws1.cell(row=idx, column=2, value=req_text).font = Font(name="Arial", size=10)
        ws1.cell(row=idx, column=3, value="")  # Compliance cell
        ws1.cell(row=idx, column=4, value="")  # Remarks cell

    # Row 6: Summary formula
    ws1.cell(row=6, column=1, value="Total Items")
    ws1.cell(row=6, column=2, value="=COUNT(A2:A5)")

    # Merged cell in row 7
    ws1.merge_cells("A7:D7")
    ws1["A7"] = "Confidential - Vendor RFP Response"

    # Hidden column E
    ws1.column_dimensions["E"].hidden = True

    # Sheet 2: Network Specs
    ws2 = wb.create_sheet(title="Network Specs")
    for col_idx, h in enumerate(headers, start=1):
        ws2.cell(row=1, column=col_idx, value=h)
    ws2.cell(row=2, column=1, value=1)
    ws2.cell(row=2, column=2, value="4x 25GbE SFP28 network ports")
    ws2.cell(row=2, column=3, value="")
    ws2.cell(row=2, column=4, value="")

    file_path = tmp_path / "sample_rfp_template.xlsx"
    wb.save(file_path)
    return file_path


def test_excel_writer_populates_workbook_and_preserves_formatting(
    sample_rfp_workbook: Path, tmp_path: Path
) -> None:
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(sample_rfp_workbook), workbook_id="test-wb-1", version=1)

    decisions = [
        ComplianceDecision(
            requirement_text="Minimum 256GB DDR5 ECC RAM required",
            state=ComplianceState.COMPLIANT,
            confidence=0.95,
            reasoning="Specification satisfies threshold.",
            resolving_layer="rule_based",
            evidence=[
                EvidenceItem(
                    source_page=4,
                    source_table_id="T4-1",
                    value="512GB DDR5",
                    confidence=0.95,
                    extraction_method="rule_based",
                    citation="Page 4, Table T4-1",
                    reasoning="512GB meets required minimum of 256GB",
                )
            ],
        ),
        ComplianceDecision(
            requirement_text="Dual redundant 1600W Titanium power supplies",
            state=ComplianceState.COMPLIANT,
            confidence=0.92,
            reasoning="Dual 1600W hot-plug PSUs confirmed.",
            resolving_layer="rule_based",
            evidence=[
                EvidenceItem(
                    source_page=2,
                    value="Dual 1600W Titanium PSUs",
                    confidence=0.92,
                    extraction_method="rule_based",
                    citation="Page 2",
                    reasoning="Redundant power supplies present.",
                )
            ],
        ),
        ComplianceDecision(
            requirement_text="At least 100TB All-Flash NVMe usable storage",
            state=ComplianceState.COMPLIANT,
            confidence=0.94,
            reasoning="150TB NVMe storage meets requirement.",
            resolving_layer="unit_conversion",
            evidence=[
                EvidenceItem(
                    source_page=5,
                    value="150TB All-Flash NVMe",
                    confidence=0.94,
                    extraction_method="unit_conversion",
                    citation="Page 5",
                    reasoning="150TB > 100TB",
                )
            ],
        ),
        ComplianceDecision(
            requirement_text="Liquid immersion cooling system",
            state=ComplianceState.NOT_FOUND,
            confidence=0.90,
            reasoning="Specification not found in reference data.",
            resolving_layer="none",
            evidence=[],
        ),
        ComplianceDecision(
            requirement_text="4x 25GbE SFP28 network ports",
            state=ComplianceState.COMPLIANT,
            confidence=0.88,
            reasoning="Dual-port 25GbE adapters present.",
            resolving_layer="semantic_match",
            evidence=[
                EvidenceItem(
                    source_page=6,
                    value="4x 25GbE SFP28",
                    confidence=0.88,
                    extraction_method="semantic_match",
                    citation="Page 6",
                    reasoning="Network adapters meet spec.",
                )
            ],
        ),
    ]

    writer = ExcelWriter(output_dir=tmp_path / "output")
    result = writer.populate_workbook(
        template_path=sample_rfp_workbook,
        analysis=analysis,
        decisions=decisions,
    )

    # 1. Output file assertions
    assert Path(result.output_file_path).exists()
    assert "_populated_" in result.filename
    assert result.total_populated_cells > 0
    assert result.summary_sheet_created is True

    # 2. Open populated workbook and check preservation
    pop_wb = openpyxl.load_workbook(result.output_file_path, data_only=False)

    # Check Summary Sheet exists as Sheet 0
    assert pop_wb.sheetnames[0] == "Compliance Summary"
    sum_ws = pop_wb["Compliance Summary"]
    assert sum_ws["B2"].value == "RFP COMPLIANCE EVALUATION REPORT"

    # Check Original Sheets exist
    assert "Server Specs" in pop_wb.sheetnames
    assert "Network Specs" in pop_wb.sheetnames

    ws1 = pop_wb["Server Specs"]
    # Check compliance values populated
    assert ws1["C2"].value == "Compliant"
    assert "Page 4, Table T4-1" in str(ws1["D2"].value)
    assert ws1["C5"].value == "Not Found"

    # Check formula preserved
    assert ws1["B6"].value == "=COUNT(A2:A5)"

    # Check merged cells preserved
    assert "A7:D7" in [str(r) for r in ws1.merged_cells.ranges]
    assert ws1["A7"].value == "Confidential - Vendor RFP Response"

    # Check hidden column preserved
    assert ws1.column_dimensions["E"].hidden is True

    # 3. Validation report assertions
    assert result.validation_report.is_valid is True
    assert result.validation_report.passed_checks >= 5


def test_excel_validator_flags_low_confidence_and_missing_evidence(
    sample_rfp_workbook: Path, tmp_path: Path
) -> None:
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(sample_rfp_workbook), workbook_id="test-wb-flags", version=1)

    low_conf_decisions = [
        ComplianceDecision(
            requirement_text="Minimum 256GB DDR5 ECC RAM required",
            state=ComplianceState.PARTIALLY_COMPLIANT,
            confidence=0.55,  # Low confidence < 0.70
            reasoning="Uncertain memory spec.",
            resolving_layer="semantic_match",
            evidence=[],  # Missing evidence
        )
    ]

    writer = ExcelWriter(output_dir=tmp_path / "output_flags")
    result = writer.populate_workbook(
        template_path=sample_rfp_workbook,
        analysis=analysis,
        decisions=low_conf_decisions,
    )

    report = result.validation_report
    assert report.is_valid is True
    assert any(issue.rule == "low_confidence_flagged" for issue in report.issues)


def test_excel_validator_detects_corrupt_or_missing_sheets(
    sample_rfp_workbook: Path, tmp_path: Path
) -> None:
    analyzer = ExcelAnalyzer()
    analysis = analyzer.analyze_workbook(str(sample_rfp_workbook), workbook_id="test-wb-2", version=1)

    validator = ExcelValidator()

    # Test with non-existent file
    report_missing = validator.validate(tmp_path / "does_not_exist.xlsx", analysis, [])
    assert report_missing.is_valid is False
    assert any(issue.rule == "file_existence" for issue in report_missing.issues)

    # Test with workbook missing original sheet
    wb_broken = openpyxl.Workbook()
    wb_broken.active.title = "Some Other Sheet"
    broken_file = tmp_path / "broken.xlsx"
    wb_broken.save(broken_file)

    report_broken = validator.validate(broken_file, analysis, [])
    assert report_broken.is_valid is False
    assert any(issue.rule == "sheet_preservation" for issue in report_broken.issues)
